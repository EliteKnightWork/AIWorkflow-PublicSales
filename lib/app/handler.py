import base64
import io
import os
import json
import logging
import pdfplumber
import requests
import re
from service.anthropic_service import AnthropicService
from pptx import Presentation
import docx2txt

logger = logging.getLogger()
logger.setLevel(logging.INFO)

def create_proposal(event, context):
    logger.info(f"Event: {json.dumps(event)}")
    
    for record in event['Records']:
        api_endpoint_url = os.environ.get('API_ENDPOINT_URL')
        message = json.loads(record['body'])

        file_text=""

        template_url=None if 'template_presigned_url' not in message else message['template_presigned_url']
        project_info = message['project_info']
        company_name = message['company_name']
        job_message_id = record['messageId']
        
        # ファイルをダウンロード    
        for doc_url in project_info["external_document_urls"]:
            logger.info('doc_url: %s', doc_url)

            if doc_url.startswith('http://localhost'):
                doc_url = doc_url.replace('http://localhost', 'http://app')

            # ファイルをダウンロード
            response = requests.get(doc_url)

            if response.status_code == 200:
                # ファイルをバイナリとして取得
                file_content = io.BytesIO(response.content)
                
                try:
                    # Content-Typeヘッダーからファイルタイプを判定
                    content_type = response.headers.get('Content-Type', '').lower()
                    
                    if 'application/pdf' in content_type:
                        # PDFファイルの場合
                        with pdfplumber.open(file_content) as pdf:
                            # 全ページのテキストを抽出
                            for page in pdf.pages:
                                extracted_text = page.extract_text()
                                if extracted_text:
                                    file_text += extracted_text + "\n"
                        
                        logger.info(f'PDFからテキストを抽出しました: {len(file_text)} 文字')
                    elif 'application/vnd.openxmlformats-officedocument.wordprocessingml.document' in content_type or 'application/msword' in content_type:
                        # Wordファイルの場合
                        # docx2txtを使用してWordファイルからテキストを抽出
                        text = docx2txt.process(file_content)
                        if text:
                            file_text += text + "\n"
                            logger.info(f'docx2txtを使用してWordファイルからテキストを抽出しました: {len(text)} 文字')
                        else:
                            logger.warning('docx2txtでテキスト抽出できませんでした')
                            
                        
                        logger.info(f'Wordからテキストを抽出しました: {len(file_text)} 文字')
                    else:
                        # Content-Typeが不明確な場合は拡張子でバックアップ判定
                        file_extension = os.path.splitext(doc_url)[1].lower()
                        if file_extension == '.pdf':
                            with pdfplumber.open(file_content) as pdf:
                                for page in pdf.pages:
                                    extracted_text = page.extract_text()
                                    if extracted_text:
                                        file_text += extracted_text + "\n"
                            logger.info(f'拡張子からPDFと判断してテキストを抽出しました: {len(file_text)} 文字')
                        elif file_extension in ['.doc', '.docx']:
                            # docx2txtを使用してWordファイルからテキストを抽出
                            text = docx2txt.process(file_content)
                            if text:
                                file_text += text + "\n"
                                logger.info(f'docx2txtを使用してWordファイルからテキストを抽出しました: {len(text)} 文字')
                            else:
                                logger.warning('docx2txtでテキスト抽出できませんでした')
                                
                        else:
                            logger.warning(f'サポートされていないファイル形式です: {content_type}, 拡張子: {file_extension}')
                except Exception as e:
                    logger.error(f'テキスト抽出エラー: {str(e)}')
            else:
                logger.error(f'添付資料のダウンロードに失敗しました: {response.status_code} doc_url: {doc_url}')
                response = requests.put(f'{api_endpoint_url}/v1/proposal_jobs/{job_message_id}', json={
                    'message_id': job_message_id,
                    'is_success': False,
                    'error_message': '添付資料のダウンロードに失敗しました'
                })

                return {
                    'statusCode': 200,
                    'body': json.dumps({
                        'message': 'PDFのダウンロードに失敗しました'
                    })
                }

        try:
            anthropic_service = AnthropicService()
            outline_json = json.loads(anthropic_service.extract_proposal_outline(file_text))
            print(outline_json)
        except Exception as e:
            logger.error(f'提案書のアウトライン抽出に失敗しました: {str(e)}')
            response = requests.put(f'{api_endpoint_url}/v1/proposal_jobs/{job_message_id}', json={
                'message_id': job_message_id,
                'is_success': False,
                'error_message': '提案書のアウトライン抽出に失敗しました'
            })

            if response.status_code == 200:
                return {
                    'statusCode': 200,
                    'body': json.dumps({
                        'message': '提案書のアウトライン抽出に失敗しました'
                    })        
                }
            else:
                return {
                    'statusCode': 500,
                    'body': json.dumps({
                        'message': f'提案書のアウトライン抽出失敗通知ができませんでした code: {response.status_code}'
                    })
                }

        # テンプレートファイルを取得
        try:
            
            if template_url is None:
                logger.warning(f'テンプレートファイルが指定されていません。デフォルトのテンプレートファイルを使用します')
                # デフォルトのテンプレートファイルを使用
                template_content = _get_default_template_file()
            else:
                # get template file from s3
                logger.info(f'template_urlからファイルを取得します: {template_url}')
            
            
                template_response = requests.get(template_url)
            
                if template_response.status_code == 200:
                    # テンプレートファイルをバイナリとして取得
                    template_content = io.BytesIO(template_response.content)
                else:
                    logger.error(f'テンプレートファイルの取得に失敗しました: {template_response.status_code}')
                    logger.error(f'デフォルトのテンプレートファイルを使用します')
                    template_content = _get_default_template_file()
                
            # プレゼンテーションを作成
            prs = Presentation(template_content)
            
            # 各スライドのテキストプレースホルダーを置き換え
            for slide in prs.slides:
                # シェイプ内のテキスト置換
                for shape in slide.shapes:
                    if not shape.has_text_frame:
                        continue
                        
                    text_frame = shape.text_frame
                    for paragraph in text_frame.paragraphs:
                        
                        # 段落全体のテキストを取得
                        paragraph_text = paragraph.text
                        
                        # 基本的なプレースホルダー置換
                        if '{会社名}' in paragraph_text:
                            paragraph_text = paragraph_text.replace('{会社名}', company_name)
                        if '{タイトル}' in paragraph_text:
                            paragraph_text = paragraph_text.replace('{タイトル}', project_info.get('title', ''))
                        if '{発注機関名}' in paragraph_text:
                            paragraph_text = paragraph_text.replace('{発注機関名}', project_info.get('organization_name', ''))
                    
                        # {n_category} と {n_content} の置換
                        for num in range(1, 21):
                            if f'{{{num}_category}}' in paragraph_text:
                                if str(num) in outline_json["items"]:
                                    paragraph_text = paragraph_text.replace(f'{{{num}_category}}', outline_json["items"][str(num)]['category'])
                            if f'{{{num}_content}}' in paragraph_text:
                                if str(num) in outline_json["items"]:
                                    content = outline_json["items"][str(num)]['content']
                                    if isinstance(content, list):
                                        content_text = '\n'.join([f'- {item}' for item in content])
                                    else:
                                        content_text = content
                                    
                                    paragraph_text = paragraph_text.replace(f'{{{num}_content}}', content_text)
                            
                            if f'{{{num}_evaluation}}' in paragraph_text:
                                if str(num) in outline_json["items"]:
                                    content = outline_json["items"][str(num)]['evaluation']
                                    if isinstance(content, list):
                                        content_text = '\n'.join([f'- {item}' for item in content])
                                    else:
                                        content_text = content

                                    paragraph_text = paragraph_text.replace(f'{{{num}_evaluation}}', content_text)
                            
                            if f'{{{num}_ideas}}' in paragraph_text:
                                if str(num) in outline_json["items"]:
                                    content = outline_json["items"][str(num)]['ideas']
                                    if isinstance(content, list):
                                        content_text = '\n'.join([f'- {item}' for item in content])
                                    else:
                                        content_text = content

                                    paragraph_text = paragraph_text.replace(f'{{{num}_ideas}}', content_text)
                        
                        # 置換後のテキストを段落に設定
                        if paragraph.runs:
                            first_run = paragraph.runs[0]
                            first_run.text = paragraph_text
                            for run in paragraph.runs[1:]:
                                run.text = ""
                        else:
                            paragraph.text = paragraph_text
                
                
                # テーブル内のテキスト置換
                for shape in slide.shapes:
                    if not shape.has_table:
                        continue

                    table = shape.table
                    for row in table.rows:
                        for cell in row.cells:
                            if not cell.text_frame:
                                continue
                                
                            for paragraph in cell.text_frame.paragraphs:
                        
                                # 段落全体のテキストを取得
                                paragraph_text = paragraph.text
                                
                                # 基本的なプレースホルダー置換
                                if '{会社名}' in paragraph_text:
                                    paragraph_text = paragraph_text.replace('{会社名}', company_name)
                                if '{タイトル}' in paragraph_text:
                                    paragraph_text = paragraph_text.replace('{タイトル}', project_info.get('title', ''))
                                if '{発注機関名}' in paragraph_text:
                                    paragraph_text = paragraph_text.replace('{発注機関名}', project_info.get('organization_name', ''))
                            
                                # {n_category} と {n_content} の置換
                                for num in range(1, 21):
                                    if f'{{{num}_category}}' in paragraph_text:
                                        if str(num) in outline_json["items"]:
                                            paragraph_text = paragraph_text.replace(f'{{{num}_category}}', outline_json["items"][str(num)]['category'])
                                    if f'{{{num}_content}}' in paragraph_text:
                                        if str(num) in outline_json["items"]:
                                            content = outline_json["items"][str(num)]['content']
                                            if isinstance(content, list):
                                                content_text = '\n'.join([f'- {item}' for item in content])
                                            else:
                                                content_text = content
                                            
                                            paragraph_text = paragraph_text.replace(f'{{{num}_content}}', content_text)
                                    
                                    if f'{{{num}_evaluation}}' in paragraph_text:
                                        if str(num) in outline_json["items"]:
                                            content = outline_json["items"][str(num)]['evaluation']
                                            if isinstance(content, list):
                                                content_text = '\n'.join([f'- {item}' for item in content])
                                            else:
                                                content_text = content

                                            paragraph_text = paragraph_text.replace(f'{{{num}_evaluation}}', content_text)
                                    
                                    if f'{{{num}_ideas}}' in paragraph_text:
                                        if str(num) in outline_json["items"]:
                                            content = outline_json["items"][str(num)]['ideas']
                                            if isinstance(content, list):
                                                content_text = '\n'.join([f'- {item}' for item in content])
                                            else:
                                                content_text = content

                                            paragraph_text = paragraph_text.replace(f'{{{num}_ideas}}', content_text)
                                
                                # 置換後のテキストを段落に設定
                                if paragraph.runs:
                                    first_run = paragraph.runs[0]
                                    first_run.text = paragraph_text
                                    for run in paragraph.runs[1:]:
                                        run.text = ""
                                else:
                                    paragraph.text = paragraph_text
            
            # ファイルを保存
            output_buffer = io.BytesIO()
            prs.save(output_buffer)
            output_buffer.seek(0)
            
            # ファイルデータをblobとして準備
            project_name = project_info.get('title', 'untitled_project')
            safe_project_name = re.sub(r'[\\/*?:"<>|]', "_", project_name)
            output_filename = f"{safe_project_name}.pptx"
            
            # ファイルデータをbase64エンコード
            file_data = base64.b64encode(output_buffer.getvalue()).decode('utf-8')
            
            logger.info(f'{api_endpoint_url}/v1/proposal_jobs/{job_message_id}')
            response = requests.put(f'{api_endpoint_url}/v1/proposal_jobs/{job_message_id}', json={
                'message_id': job_message_id,
                'is_success': True,
                'proposal_file_name': output_filename,
                'proposal_content_type': 'application/vnd.openxmlformats-officedocument.presentationml.presentation',
                'proposal_file_data': file_data
            })
            logger.info(f'APIレスポンス: {response}')
            
        except Exception as e:
            logger.error(f'PowerPointファイル生成エラー: {str(e)}')
            response = requests.put(f'{api_endpoint_url}/v1/proposal_jobs/{job_message_id}', json={
                'message_id': job_message_id,
                'is_success': False,
                'error_message': '提案書のPowerPointファイル生成に失敗しました'
            })
            return {
                    'statusCode': 200,
                    'body': json.dumps({
                    'message': '提案書のPowerPointファイル生成に失敗しました'
                })
            }
        
    return {
        'statusCode': 200,
        'body': json.dumps({
            'message': '提案書の生成が完了しました'
        })
    }

def _get_default_template_file():
    template_path = os.path.join(os.path.dirname(__file__), 'slide_template_default.pptx')
    with open(template_path, 'rb') as f:
        template_content = io.BytesIO(f.read())
    return template_content


def notify_failed_proposal(event, context):
    logger.info(f"Event: {json.dumps(event)}")

    api_endpoint_url = os.environ.get('API_ENDPOINT_URL')
    
    for record in event['Records']:
        job_message_id = record['messageId']

    response = requests.put(f'{api_endpoint_url}/v1/proposal_jobs/{job_message_id}', json={
        'message_id': job_message_id,
        'is_success': False,
        'error_message': '提案書の作成中に想定外エラーが発生しました'
    })
    logger.info(f'APIレスポンス: {response}')

    return {
        'statusCode': 200,
        'body': json.dumps({
            'message': '提案書の作成に失敗しました'
        })
    }